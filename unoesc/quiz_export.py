"""
Exportação de questionários Moodle ON para DOCX e PDF.

Extrai questões do banco de questões (incluindo questionários com slots aleatórios),
obtém enunciados, alternativas e gabarito via prévia do professor, e gera arquivos
formatados para impressão ou revisão offline.
"""
from __future__ import annotations

import io
import json
import os
import re
import tempfile
import urllib.parse
from html import unescape

from bs4 import BeautifulSoup, NavigableString
from requests import Session

from .moodle import MOODLE_BASE, open_course, find_activity, find_section

_PREVIEW_OPTS = {
    "correctness": "1",
    "rightanswer": "1",
    "feedback": "1",
    "generalfeedback": "1",
    "marks": "2",
}


def _soup(html: str) -> BeautifulSoup:
    return BeautifulSoup(html, "html.parser")


def _category_filter(category_id: str | int) -> str:
    filt = {
        "category": {
            "name": "category",
            "jointype": 1,
            "values": [int(category_id)],
            "filteroptions": {"includesubcategories": True},
        }
    }
    return json.dumps(filt, separators=(",", ":"))


def _slug(texto: str) -> str:
    s = re.sub(r"\W+", "_", texto).strip("_").lower()
    return s[:80] or "questionario"


def _html_to_text(node) -> str:
    """Converte nó HTML em texto preservando quebras de parágrafo."""
    if node is None:
        return ""
    partes: list[str] = []

    def walk(el):
        if isinstance(el, NavigableString):
            partes.append(str(el))
            return
        if el.name in ("br",):
            partes.append("\n")
            return
        if el.name in ("p", "div", "li", "tr"):
            if partes and not partes[-1].endswith("\n"):
                partes.append("\n")
        for child in el.children:
            walk(child)
        if el.name in ("p", "div", "li"):
            if partes and not partes[-1].endswith("\n"):
                partes.append("\n")

    walk(node)
    texto = unescape("".join(partes))
    return re.sub(r"\n{3,}", "\n\n", texto).strip()


def _extract_images(node, session: Session) -> list[dict]:
    """Baixa imagens referenciadas no enunciado/alternativas."""
    imagens = []
    if node is None:
        return imagens
    for img in node.find_all("img"):
        src = img.get("src", "")
        if not src.startswith("http"):
            continue
        alt = img.get("alt", "imagem")
        try:
            resp = session.get(src, timeout=30)
            if resp.status_code != 200:
                continue
            ext = ".png"
            ctype = resp.headers.get("Content-Type", "")
            if "jpeg" in ctype or "jpg" in ctype:
                ext = ".jpg"
            elif "gif" in ctype:
                ext = ".gif"
            elif "svg" in ctype:
                ext = ".svg"
            imagens.append({
                "alt": alt,
                "url": src,
                "bytes": resp.content,
                "ext": ext,
            })
        except Exception:
            continue
    return imagens


def _parse_preview_form(form) -> tuple[str, dict]:
    action = form["action"].replace("&amp;", "&")
    data: dict[str, str] = {}
    for inp in form.find_all("input"):
        name = inp.get("name")
        if not name or inp.get("type") == "radio":
            continue
        data[name] = inp.get("value", "")
    return action, data


def _parse_question_element(que, session: Session) -> dict:
    """Parseia um elemento .que do Moodle."""
    qtext = que.select_one(".qtext")
    enunciado = _html_to_text(qtext)
    imagens = _extract_images(qtext, session)

    alternativas = []
    for div in que.select(".answer > div"):
        label = div.select_one("[data-region='answer-label'], .flex-fill")
        letra_el = div.select_one(".answernumber")
        texto = label.get_text(strip=True) if label else div.get_text(strip=True)
        if letra_el:
            letra = letra_el.get_text(strip=True).rstrip(".")
            if texto.startswith(letra):
                texto = texto[len(letra):].lstrip(". ")
        else:
            letra = ""
        classes = div.get("class", [])
        correta = (
            "correct" in classes
            or div.select_one(".icon.fa-check, .icon.fa-circle-check, .icon.fa-check-circle")
            is not None
        )
        alternativas.append({"letra": letra, "texto": texto, "correta": correta})

    resposta_correta = next((a for a in alternativas if a["correta"]), None)
    ra = que.select_one(".rightanswer")
    if ra and not resposta_correta:
        resposta_correta = {"texto": ra.get_text(strip=True)}

    feedback_parts = []
    for sel in (".feedback", ".specificfeedback", ".generalfeedback"):
        for el in que.select(sel):
            t = el.get_text(strip=True)
            if t:
                feedback_parts.append(t)

    tipo = "desconhecido"
    if "multichoice" in que.get("class", []):
        tipo = "multichoice"

    return {
        "enunciado": enunciado,
        "alternativas": alternativas,
        "resposta_correta": resposta_correta["texto"] if resposta_correta else "",
        "feedback": "\n".join(feedback_parts),
        "imagens": imagens,
        "tipo": tipo,
    }


def _parse_question_preview(html: str, session: Session) -> dict:
    soup = _soup(html)
    que = soup.select_one(".que")
    if not que:
        raise RuntimeError("Questão não encontrada na prévia.")
    return _parse_question_element(que, session)


def _question_id_from_html(html: str) -> str | None:
    m = re.search(r"/question/questiontext/\d+/\d+/(\d+)/", html)
    return m.group(1) if m else None


def _form_hidden_data(form) -> dict[str, str]:
    data: dict[str, str] = {}
    for inp in form.find_all("input"):
        name = inp.get("name")
        if not name:
            continue
        itype = inp.get("type")
        if itype in ("radio", "checkbox", "submit", "button"):
            continue
        data[name] = inp.get("value", "")
    return data


def _apply_answer_key(dados: dict, gabarito: dict) -> None:
    """Mescla gabarito do banco na questão extraída da tentativa."""
    rc = gabarito.get("resposta_correta", "")
    if rc:
        dados["resposta_correta"] = rc
        rc_norm = rc.strip().lower()
        for alt in dados.get("alternativas", []):
            txt = alt.get("texto", "").strip().lower()
            alt["correta"] = txt == rc_norm or txt in rc_norm or rc_norm in txt
    if gabarito.get("feedback"):
        dados["feedback"] = gabarito["feedback"]
    if gabarito.get("nome"):
        dados["nome"] = gabarito["nome"]


def _start_preview_attempt(session: Session, cmid: str):
    """Inicia ou continua uma prévia do questionário como professor."""
    resp = session.get(f"{MOODLE_BASE}/mod/quiz/view.php", params={"id": cmid})
    form_start = _soup(resp.text).find("form", action=re.compile("startattempt"))
    if not form_start:
        raise RuntimeError("Prévia do questionário indisponível (sem permissão ou quiz inexistente).")

    data = {
        inp["name"]: inp.get("value", "")
        for inp in form_start.find_all("input")
        if inp.get("name")
    }
    return session.post(
        form_start["action"],
        data=data,
        allow_redirects=True,
    )


def _attempt_id(resp) -> str | None:
    m = re.search(r"attempt=(\d+)", resp.url) or re.search(r"attempt=(\d+)", resp.text)
    return m.group(1) if m else None


def extract_quiz_questions(
    session: Session,
    cmid: str,
    course_id: str,
    include_answer_key: bool = True,
    max_questoes: int | None = None,
) -> list[dict]:
    """Extrai as questões sorteadas em uma prévia do questionário (não do banco).

    Inicia 'Ver prévia do questionário' como professor e percorre todas as
    páginas da tentativa, coletando exatamente as questões que compõem o quiz.
    """
    resp = _start_preview_attempt(session, cmid)
    attempt = _attempt_id(resp)
    if not attempt:
        raise RuntimeError("Não foi possível obter o ID da tentativa de prévia.")

    questoes: list[dict] = []
    vistos: set[str] = set()
    numero = 0

    for page in range(50):
        rp = session.get(
            f"{MOODLE_BASE}/mod/quiz/attempt.php",
            params={"attempt": attempt, "cmid": cmid, "page": page},
        )
        soup = _soup(rp.text)
        pagina_tem_questao = False

        for que in soup.select(".que"):
            qtext = que.select_one(".qtext")
            if not qtext or not qtext.get_text(strip=True):
                continue

            que_html = str(que)
            chave = _question_id_from_html(que_html) or _html_to_text(qtext)[:200]
            if chave in vistos:
                continue
            vistos.add(chave)
            pagina_tem_questao = True
            numero += 1

            dados = _parse_question_element(que, session)
            dados["numero"] = numero
            qid = _question_id_from_html(que_html)
            if qid:
                dados["id"] = qid
                if include_answer_key:
                    try:
                        gab = get_question(session, course_id, qid)
                        _apply_answer_key(dados, gab)
                    except Exception:
                        pass
            dados.setdefault("nome", f"Questão {numero}")
            questoes.append(dados)

            if max_questoes and len(questoes) >= max_questoes:
                return questoes

        if not pagina_tem_questao:
            break

    if not questoes:
        raise RuntimeError("Nenhuma questão encontrada na prévia do questionário.")

    return questoes


def list_category_questions(
    session: Session,
    course_id: str,
    category_id: str | int | None = None,
    filter_param: str | None = None,
) -> list[dict]:
    """Lista questões de uma categoria do banco de questões.

    Args:
        category_id:   ID da categoria (alternativa a filter_param).
        filter_param:  parâmetro filter já codificado (copiado do link do Moodle).

    Returns:
        list[dict]: [{id, nome, tipo}, ...]
    """
    if filter_param:
        filt = filter_param
    elif category_id is not None:
        filt = _category_filter(category_id)
    else:
        raise ValueError("Informe category_id ou filter_param.")

    if filter_param:
        # filter_param vem URL-encoded do href do Moodle — não re-encodar
        url = f"{MOODLE_BASE}/question/edit.php?courseid={course_id}&filter={filter_param}"
        resp = session.get(url)
    else:
        resp = session.get(
            f"{MOODLE_BASE}/question/edit.php",
            params={"courseid": course_id, "filter": filt},
        )
    soup = _soup(resp.text)
    questoes = []
    vistos: set[str] = set()

    for inp in soup.select('input[id^="checkq"]'):
        qid = inp.get("name", "").lstrip("q")
        if not qid or qid in vistos:
            continue
        vistos.add(qid)
        row = inp.find_parent("tr")
        nome = qid
        tipo = ""
        if row:
            name_el = row.select_one("[data-itemtype='questionname']")
            if name_el:
                nome = name_el.get("data-value") or name_el.get_text(strip=True)
            tipo_cell = row.select_one("[data-columnid*='question_type']")
            if tipo_cell:
                tipo = tipo_cell.get_text(strip=True)
        questoes.append({"id": qid, "nome": nome, "tipo": tipo})

    return questoes


def get_question(session: Session, course_id: str, question_id: str) -> dict:
    """Obtém enunciado, alternativas e gabarito de uma questão via prévia do professor."""
    params = {
        "id": question_id,
        "restartversion": "0",
        "courseid": course_id,
        **_PREVIEW_OPTS,
    }
    resp = session.get(f"{MOODLE_BASE}/question/bank/previewquestion/preview.php", params=params)
    soup = _soup(resp.text)
    form = soup.find("form", id="responseform")
    if not form:
        dados = _parse_question_preview(resp.text, session)
        return {"id": question_id, **dados}

    action, data = _parse_preview_form(form)
    data["fill"] = "Preencher com respostas corretas"
    resp2 = session.post(action, data=data)
    dados = _parse_question_preview(resp2.text, session)
    return {"id": question_id, **dados}


def list_question_bank_categories(session: Session, course_id: str) -> list[dict]:
    """Lista categorias do banco de questões de um curso Moodle.

    Returns:
        list[dict]: [{id, nome}, ...] — apenas categorias do curso (IDs altos).
    """
    resp = session.get(
        f"{MOODLE_BASE}/question/edit.php",
        params={"courseid": course_id},
    )
    soup = _soup(resp.text)
    categorias: list[dict] = []
    vistos: set[str] = set()

    for opt in soup.select("option"):
        val = opt.get("value", "")
        if not val.isdigit() or int(val) < 10000:
            continue
        if val in vistos:
            continue
        nome = opt.get_text(strip=True)
        if not nome:
            continue
        vistos.add(val)
        categorias.append({"id": val, "nome": nome})

    return categorias


def extract_question_bank(
    session: Session,
    course_id: str,
    title: str | None = None,
    categoria_ids: list[str] | None = None,
    include_answer_key: bool = True,
) -> dict:
    """Extrai todas as questões do banco de questões de um curso."""
    categorias = list_question_bank_categories(session, course_id)
    if categoria_ids:
        ids_set = {str(i) for i in categoria_ids}
        categorias = [c for c in categorias if c["id"] in ids_set]
    if not categorias:
        raise RuntimeError(f"Nenhuma categoria encontrada no banco do curso {course_id}.")

    if title is None:
        resp = session.get(f"{MOODLE_BASE}/course/view.php", params={"id": course_id})
        m = re.search(r"<title>(.*?)</title>", resp.text, re.I)
        title = m.group(1).strip() if m else f"curso_{course_id}"
        title = re.sub(r"\s*\|\s*On$", "", title)

    blocos = []
    total = 0
    vistos_global: set[str] = set()

    for cat in categorias:
        print(f"  Categoria [{cat['id']}] {cat['nome']}...")
        lista = list_category_questions(session, course_id, cat["id"])
        questoes = []
        for q in lista:
            if q["id"] in vistos_global:
                continue
            vistos_global.add(q["id"])
            print(f"    Extraindo [{q['id']}] {q['nome']}...")
            try:
                dados = get_question(session, course_id, q["id"])
                dados["nome"] = q["nome"]
                questoes.append(dados)
            except Exception as exc:
                questoes.append({
                    "id": q["id"],
                    "nome": q["nome"],
                    "enunciado": f"[Erro ao extrair questão: {exc}]",
                    "alternativas": [],
                    "resposta_correta": "",
                    "feedback": "",
                    "imagens": [],
                    "tipo": q.get("tipo", ""),
                })
        total += len(questoes)
        if questoes:
            blocos.append({
                "category_id": cat["id"],
                "categoria": cat["nome"],
                "questoes": questoes,
            })

    if total == 0:
        raise RuntimeError("Nenhuma questão encontrada no banco de questões.")

    return {
        "title": title,
        "course_id": course_id,
        "total_questoes": total,
        "source": "banco",
        "blocos": blocos,
        "include_answer_key": include_answer_key,
    }


def export_question_bank(
    session: Session,
    course_id: str,
    title: str | None = None,
    categoria_ids: list[str] | None = None,
    destination: str | None = None,
    formats: tuple[str, ...] = ("docx", "pdf"),
    include_answer_key: bool = True,
) -> dict:
    """Exporta banco de questões de um curso para DOCX/PDF."""
    print(f"Extraindo banco de questões (course_id={course_id})...")
    dados = extract_question_bank(
        session, course_id, title=title,
        categoria_ids=categoria_ids,
        include_answer_key=include_answer_key,
    )

    base = destination
    if base is None:
        base = os.path.expanduser(f"~/Downloads/{_slug(dados['title'])}_banco")
    elif base.endswith((".docx", ".pdf")):
        base = os.path.splitext(base)[0]

    resultado = {"dados": dados, "docx": None, "pdf": None}
    if "docx" in formats:
        resultado["docx"] = export_quiz_docx(dados, f"{base}.docx")
    if "pdf" in formats:
        resultado["pdf"] = export_quiz_pdf(dados, f"{base}.pdf")
    return resultado


def export_discipline_question_bank(
    session: Session,
    discipline_query: str,
    destination: str | None = None,
    formats: tuple[str, ...] = ("docx", "pdf"),
    include_answer_key: bool = True,
    categoria_ids: list[str] | None = None,
) -> dict:
    """Fluxo completo: discipline → banco de questões → exportação."""
    from .portal import find_discipline

    disc = find_discipline(session, discipline_query)
    if not disc:
        raise ValueError(f"Disciplina '{discipline_query}' não encontrada.")

    session, course_id = open_course(session, disc["dof"])
    title = f"{disc['codigo']} - {disc['nome']}"

    print(f"Disciplina: {title}")
    print(f"Course ID: {course_id}")

    resultado = export_question_bank(
        session, course_id,
        title=title,
        categoria_ids=categoria_ids,
        destination=destination,
        formats=formats,
        include_answer_key=include_answer_key,
    )
    resultado["discipline"] = disc
    return resultado


def analyze_quiz(session: Session, cmid: str) -> dict:
    """Analisa a estrutura de um questionário (edit.php).

    Returns:
        dict com {title, cmid, course_id, slots, categorias}
    """
    resp = session.get(f"{MOODLE_BASE}/mod/quiz/edit.php", params={"cmid": cmid})
    soup = _soup(resp.text)
    titulo_el = soup.find("title")
    title = titulo_el.get_text(strip=True) if titulo_el else ""
    title = re.sub(r"^Editando questionário:\s*", "", title)
    title = re.sub(r"\s*\|\s*On$", "", title)

    m_course = re.search(r'"courseId"\s*:\s*(\d+)', resp.text)
    course_id = m_course.group(1) if m_course else None

    slots = []
    categorias: dict[str, str] = {}

    for li in soup.select("li.activity.slot"):
        slot_id = li.get("id", "").replace("slot-", "")
        classes = li.get("class", [])
        tipo_slot = "random" if "random" in classes else "fixed"
        nome_el = li.select_one(".questionname")
        nome = nome_el.get_text(strip=True) if nome_el else ""

        cat_link = li.select_one("a.mod_quiz_random_qbank_link")
        cat_id = None
        cat_filter = None
        if cat_link and cat_link.get("href"):
            href = cat_link["href"]
            m = re.search(r"values%22%3A%5B(\d+)%5D", href)
            if m:
                cat_id = m.group(1)
            mf = re.search(r"filter=([^&]+)", href)
            if mf:
                cat_filter = mf.group(1)
                m_cat = re.search(r"\(([^)]+)\)", nome)
                cat_label = m_cat.group(1) if m_cat else nome
                cat_label = re.sub(r"\s+e subcategorias$", "", cat_label)
                categorias[cat_id or cat_filter] = {
                    "id": cat_id,
                    "filter": cat_filter,
                    "nome": cat_label,
                }

        slots.append({
            "slot_id": slot_id,
            "tipo": tipo_slot,
            "nome": nome,
            "category_id": cat_id,
        })

    return {
        "title": title,
        "cmid": cmid,
        "course_id": course_id,
        "slots": slots,
        "categorias": categorias,
    }


def extract_quiz(
    session: Session,
    course_id: str,
    cmid: str,
    include_answer_key: bool = True,
    source: str = "quiz",
) -> dict:
    """Extrai questões de um questionário Moodle.

    Args:
        source: 'quiz' = questões sorteadas na prévia do questionário (padrão);
               'banco' = todas as questões das categorias do banco (legado).

    Returns:
        dict com metadados e lista de questões.
    """
    info = analyze_quiz(session, cmid)

    if source == "quiz":
        print("Extraindo questões via prévia do questionário...")
        n_slots = len(info.get("slots") or [])
        questoes = extract_quiz_questions(
            session, cmid, course_id, include_answer_key,
            max_questoes=n_slots or None,
        )
        return {
            "title": info["title"],
            "cmid": cmid,
            "course_id": course_id,
            "slots": info["slots"],
            "total_questoes": len(questoes),
            "source": "quiz",
            "blocos": [{"categoria": "Questionário", "questoes": questoes}],
            "include_answer_key": include_answer_key,
        }

    # source == "banco" — exporta categorias inteiras do banco de questões
    info = analyze_quiz(session, cmid)
    cat_ids = list(info["categorias"].keys())
    if not cat_ids:
        raise RuntimeError(
            "Nenhuma categoria de questões encontrada no questionário. "
            "Verifique se você tem permissão de edição."
        )

    blocos = []
    total = 0
    vistos_global: set[str] = set()
    for cat_key in cat_ids:
        cat_info = info["categorias"][cat_key]
        if isinstance(cat_info, dict):
            meta = cat_info.get("nome", cat_key)
            cat_id = cat_info.get("id")
            cat_filter = cat_info.get("filter")
        else:
            meta, cat_id, cat_filter = cat_info, cat_key, None

        lista = list_category_questions(
            session, course_id,
            category_id=cat_id,
            filter_param=cat_filter,
        )
        questoes = []
        for q in lista:
            if q["id"] in vistos_global:
                continue
            vistos_global.add(q["id"])
            print(f"  Extraindo [{q['id']}] {q['nome']}...")
            try:
                dados = get_question(session, course_id, q["id"])
                dados["nome"] = q["nome"]
                questoes.append(dados)
            except Exception as exc:
                questoes.append({
                    "id": q["id"],
                    "nome": q["nome"],
                    "enunciado": f"[Erro ao extrair questão: {exc}]",
                    "alternativas": [],
                    "resposta_correta": "",
                    "feedback": "",
                    "imagens": [],
                    "tipo": q.get("tipo", ""),
                })
        total += len(questoes)
        if questoes:
            blocos.append({
                "category_id": cat_id or cat_key,
                "categoria": meta,
                "questoes": questoes,
            })

    return {
        "title": info["title"],
        "cmid": cmid,
        "course_id": course_id,
        "slots": info["slots"],
        "total_questoes": total,
        "source": "banco",
        "blocos": blocos,
        "include_answer_key": include_answer_key,
    }


def _save_temp_image(img: dict) -> str | None:
    if not img.get("bytes"):
        return None
    ext = img.get("ext", ".png")
    if ext == ".svg":
        try:
            import cairosvg
            png_bytes = cairosvg.svg2png(bytestring=img["bytes"])
            ext = ".png"
            img_bytes = png_bytes
        except Exception:
            return None
    else:
        img_bytes = img["bytes"]
    fd, path = tempfile.mkstemp(suffix=ext)
    os.close(fd)
    with open(path, "wb") as f:
        f.write(img_bytes)
    return path


def export_quiz_docx(dados: dict, destination: str | None = None) -> str:
    """Gera arquivo DOCX a partir dos dados extraídos."""
    from docx import Document
    from docx.shared import Inches, Pt

    if destination is None:
        destination = os.path.expanduser(f"~/Downloads/{_slug(dados['title'])}.docx")

    doc = Document()
    doc.add_heading(dados["title"], level=0)
    doc.add_paragraph(f"Total de questões: {dados['total_questoes']}")

    temp_files: list[str] = []
    num = 0

    try:
        for bloco in dados["blocos"]:
            if bloco.get("categoria") and bloco["categoria"] != "Questionário":
                doc.add_heading(bloco["categoria"], level=1)
            for q in bloco["questoes"]:
                num += 1
                titulo_q = q.get("nome") or f"Questão {num}"
                doc.add_heading(f"{num}. {titulo_q}", level=2)

                if q.get("enunciado"):
                    for par in q["enunciado"].split("\n"):
                        if par.strip():
                            doc.add_paragraph(par.strip())

                for img in q.get("imagens", []):
                    path = _save_temp_image(img)
                    if path:
                        temp_files.append(path)
                        try:
                            doc.add_picture(path, width=Inches(5))
                        except Exception:
                            doc.add_paragraph(f"[Imagem: {img.get('alt', '')}]")

                alts = q.get("alternativas", [])
                if alts:
                    for alt in alts:
                        prefix = alt.get("letra", "")
                        texto = alt.get("texto", "")
                        linha = f"{prefix}) {texto}" if prefix else texto
                        p = doc.add_paragraph(linha)
                        if dados.get("include_answer_key") and alt.get("correta"):
                            for run in p.runs:
                                run.bold = True

                if dados.get("include_answer_key") and q.get("resposta_correta"):
                    p = doc.add_paragraph()
                    run = p.add_run(f"Gabarito: {q['resposta_correta']}")
                    run.bold = True
                    run.font.size = Pt(10)

                if dados.get("include_answer_key") and q.get("feedback"):
                    doc.add_paragraph(f"Feedback: {q['feedback']}")

                doc.add_paragraph("")
    finally:
        for path in temp_files:
            try:
                os.remove(path)
            except OSError:
                pass

    doc.save(destination)
    print(f"DOCX salvo em: {destination}")
    return destination


def export_quiz_pdf(dados: dict, destination: str | None = None) -> str:
    """Gera arquivo PDF a partir dos dados extraídos."""
    from fpdf import FPDF

    if destination is None:
        destination = os.path.expanduser(f"~/Downloads/{_slug(dados['title'])}.pdf")

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    font_path = "/System/Library/Fonts/Supplemental/Arial Unicode.ttf"
    if not os.path.exists(font_path):
        font_path = "/Library/Fonts/Arial Unicode.ttf"
    if os.path.exists(font_path):
        pdf.add_font("ArialUni", "", font_path)
        pdf.set_font("ArialUni", size=11)
        font_name = "ArialUni"
    else:
        pdf.set_font("Helvetica", size=11)
        font_name = "Helvetica"

    def write_line(text: str, size: int = 11, bold: bool = False):
        pdf.set_font(font_name, size=size)
        pdf.multi_cell(0, 6, text)
        pdf.ln(1)

    write_line(dados["title"], size=16)
    write_line(f"Total de questões: {dados['total_questoes']}", size=10)
    pdf.ln(4)

    temp_files: list[str] = []
    num = 0

    try:
        for bloco in dados["blocos"]:
            write_line(bloco["categoria"], size=13)
            pdf.ln(2)

            for q in bloco["questoes"]:
                num += 1
                titulo_q = q.get("nome") or f"Questão {num}"
                write_line(f"{num}. {titulo_q}", size=12)

                if q.get("enunciado"):
                    write_line(q["enunciado"])

                for img in q.get("imagens", []):
                    path = _save_temp_image(img)
                    if path:
                        temp_files.append(path)
                        try:
                            pdf.image(path, w=120)
                            pdf.ln(4)
                        except Exception:
                            write_line(f"[Imagem: {img.get('alt', '')}]")

                for alt in q.get("alternativas", []):
                    prefix = alt.get("letra", "")
                    texto = alt.get("texto", "")
                    linha = f"  {prefix}) {texto}" if prefix else f"  {texto}"
                    if dados.get("include_answer_key") and alt.get("correta"):
                        linha += "  [correta]"
                    write_line(linha)

                if dados.get("include_answer_key") and q.get("resposta_correta"):
                    write_line(f"Gabarito: {q['resposta_correta']}")
                if dados.get("include_answer_key") and q.get("feedback"):
                    write_line(f"Feedback: {q['feedback']}")

                pdf.ln(4)
    finally:
        for path in temp_files:
            try:
                os.remove(path)
            except OSError:
                pass

    pdf.output(destination)
    print(f"PDF salvo em: {destination}")
    return destination


def export_quiz(
    session: Session,
    course_id: str,
    cmid: str,
    destination: str | None = None,
    formats: tuple[str, ...] = ("docx", "pdf"),
    include_answer_key: bool = True,
    source: str = "quiz",
) -> dict:
    """Extrai questionário e exporta nos formats solicitados.

    Args:
        course_id: ID do curso no Moodle.
        cmid:      ID da atividade quiz (campo activity_id de list_activities).
        destination:   Caminho base sem extensão. None = ~/Downloads/<title>.
        formats:  ('docx', 'pdf') ou subset.
        include_answer_key: destaca respostas corretas e feedback.
        source:            'quiz' (padrão) ou 'banco'.

    Returns:
        dict com {dados, docx, pdf} conforme formats gerados.
    """
    print(f"Extraindo questionário (cmid={cmid})...")
    dados = extract_quiz(
        session, course_id, cmid,
        include_answer_key=include_answer_key,
        source=source,
    )

    base = destination
    if base is None:
        base = os.path.expanduser(f"~/Downloads/{_slug(dados['title'])}")
    elif base.endswith((".docx", ".pdf")):
        base = os.path.splitext(base)[0]

    resultado = {"dados": dados, "docx": None, "pdf": None}
    if "docx" in formats:
        resultado["docx"] = export_quiz_docx(dados, f"{base}.docx")
    if "pdf" in formats:
        resultado["pdf"] = export_quiz_pdf(dados, f"{base}.pdf")
    return resultado


def export_discipline_quiz(
    session: Session,
    discipline_query: str,
    section_query: str,
    activity_query: str,
    destination: str | None = None,
    formats: tuple[str, ...] = ("docx", "pdf"),
    include_answer_key: bool = True,
    source: str = "quiz",
    open_course_fn=None,
    find_discipline_fn=None,
) -> dict:
    """Fluxo completo: discipline → seção → atividade quiz → exportação.

    Args:
        discipline_query: ex. '10276/EAD54-12' ou 'Banco de Dados'.
        section_query:      ex. 'Semana 7'.
        activity_query:  ex. 'Prova Objetiva'.
        open_course_fn:   injetável para testes (default: moodle.open_course).
        find_discipline_fn: injetável para testes.

    Returns:
        dict com paths gerados e metadados.
    """
    if open_course_fn is None:
        from . import moodle as moodle_mod
        open_course_fn = moodle_mod.open_course
    if find_discipline_fn is None:
        from . import portal as portal_mod
        find_discipline_fn = portal_mod.find_discipline

    disc = find_discipline_fn(session, discipline_query)
    if not disc:
        raise ValueError(f"Disciplina '{discipline_query}' não encontrada.")

    session, course_id = open_course_fn(session, disc["dof"])
    secao = find_section(session, course_id, section_query)
    if not secao:
        raise ValueError(f"Seção '{section_query}' não encontrada no curso {course_id}.")

    atividade = find_activity(session, secao["section_id"], activity_query, tipo="quiz")
    if not atividade:
        raise ValueError(
            f"Questionário '{activity_query}' não encontrado na seção '{secao['nome']}'. "
            "Verifique se a atividade é do tipo quiz nativo (não LTI)."
        )

    print(f"Disciplina: {disc['codigo']} - {disc['nome']}")
    print(f"Seção: {secao['nome']}")
    print(f"Atividade: {atividade['nome']} (cmid={atividade['activity_id']})")

    resultado = export_quiz(
        session,
        course_id,
        atividade["activity_id"],
        destination=destination,
        formats=formats,
        include_answer_key=include_answer_key,
        source=source,
    )
    resultado["discipline"] = disc
    resultado["secao"] = secao
    resultado["atividade"] = atividade
    return resultado
